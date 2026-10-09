from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# 프로젝트 루트/data/GreenStep_관리비_탄소계수.xlsx
PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXCEL_PATH = PROJECT_ROOT / "data" / "GreenStep_관리비_탄소계수.xlsx"

MASTER_SHEET_NAME = "관리비 탄소계수 마스터"


class UtilityCarbonService:
    """
    관리비 고지서의 사용량/금액을 기반으로 탄소배출량을 계산하는 서비스.

    계산 우선순위
    1. USAGE      : 실제 사용량 기반
    2. COST       : 세부 요금 기반 fallback
    3. COST_TOTAL : 전체 관리비 기반 최종 fallback

    특수 유형
    - SOURCE_BASED : 급탕처럼 실제 열원에 따라 계산
    - UNSUPPORTED  : 하수처럼 현재 직접 계산하지 않는 항목
    """

    def __init__(self, excel_path: Path | None = None):
        self.excel_path = excel_path or EXCEL_PATH
        self.coefficients = self._load_coefficients()

    def _load_coefficients(self) -> dict[str, dict[str, Any]]:
        """
        Excel의 '관리비 탄소계수 마스터' 시트를 읽어
        세부 항목명을 key로 하는 dictionary를 생성한다.
        """
        if not self.excel_path.exists():
            raise FileNotFoundError(
                f"관리비 탄소계수 파일을 찾을 수 없습니다: "
                f"{self.excel_path}"
            )

        workbook = load_workbook(
            filename=self.excel_path,
            read_only=True,
            data_only=True,
        )

        try:
            if MASTER_SHEET_NAME not in workbook.sheetnames:
                raise ValueError(
                    f"'{MASTER_SHEET_NAME}' 시트를 찾을 수 없습니다. "
                    f"현재 시트: {workbook.sheetnames}"
                )

            sheet = workbook[MASTER_SHEET_NAME]
            rows = sheet.iter_rows(values_only=True)

            try:
                headers = next(rows)
            except StopIteration:
                raise ValueError(
                    "탄소계수 Excel 파일이 비어 있습니다."
                )

            header_map = {
                str(header).strip(): index
                for index, header in enumerate(headers)
                if header is not None
            }

            required_columns = [
                "세부 항목",
                "계수 유형",
                "우선순위",
                "계산 기준",
                "배출계수",
                "단위",
            ]

            missing_columns = [
                column
                for column in required_columns
                if column not in header_map
            ]

            if missing_columns:
                raise ValueError(
                    "탄소계수 Excel에 필요한 열이 없습니다: "
                    f"{missing_columns}"
                )

            coefficients: dict[str, dict[str, Any]] = {}

            for row in rows:
                item_name = row[header_map["세부 항목"]]

                if not item_name:
                    continue

                coefficient = row[header_map["배출계수"]]

                coefficients[str(item_name).strip()] = {
                    "type": row[header_map["계수 유형"]],
                    "priority": row[header_map["우선순위"]],
                    "calculation_basis": row[
                        header_map["계산 기준"]
                    ],
                    "factor": (
                        float(coefficient)
                        if coefficient is not None
                        else None
                    ),
                    "unit": row[header_map["단위"]],
                    "condition": self._get_optional_value(
                        row,
                        header_map,
                        "적용 조건",
                    ),
                    "carbon_basis": self._get_optional_value(
                        row,
                        header_map,
                        "탄소 기준",
                    ),
                    "source": self._get_optional_value(
                        row,
                        header_map,
                        "출처/기준",
                    ),
                    "note": self._get_optional_value(
                        row,
                        header_map,
                        "비고",
                    ),
                }

            return coefficients

        finally:
            workbook.close()

    @staticmethod
    def _get_optional_value(
        row: tuple,
        header_map: dict[str, int],
        column_name: str,
    ) -> Any:
        """
        선택 컬럼의 값을 안전하게 가져온다.
        """
        index = header_map.get(column_name)

        if index is None:
            return None

        return row[index]

    def get_coefficient(
        self,
        item_name: str,
    ) -> dict[str, Any] | None:
        """
        항목명으로 배출계수 정보를 조회한다.

        예:
        get_coefficient("전기")
        get_coefficient("도시가스(LNG)")
        get_coefficient("전기비")
        """
        return self.coefficients.get(item_name)

    def calculate_with_factor(
        self,
        item_name: str,
        value: float,
    ) -> dict[str, Any]:
        """
        하나의 항목과 값을 전달받아 탄소배출량을 계산한다.

        USAGE / COST / COST_TOTAL 유형을 계산할 수 있다.

        SOURCE_BASED(급탕)는 별도의 열원 정보가 필요하고,
        UNSUPPORTED(하수)는 현재 계산하지 않는다.
        """
        coefficient = self.get_coefficient(item_name)

        if coefficient is None:
            return {
                "item": item_name,
                "status": "coefficient_not_found",
                "carbon_kg": None,
                "reason": (
                    "해당 항목의 탄소배출계수를 "
                    "찾을 수 없습니다."
                ),
            }

        coefficient_type = coefficient["type"]
        factor = coefficient["factor"]

        if coefficient_type == "UNSUPPORTED":
            return {
                "item": item_name,
                "status": "unsupported",
                "carbon_kg": None,
                "reason": (
                    "현재 탄소 계산 대상에서 "
                    "제외된 항목입니다."
                ),
            }

        if coefficient_type == "SOURCE_BASED":
            return {
                "item": item_name,
                "status": "source_required",
                "carbon_kg": None,
                "reason": (
                    "실제 열원 정보가 필요한 항목입니다."
                ),
            }

        if factor is None:
            return {
                "item": item_name,
                "status": "factor_missing",
                "carbon_kg": None,
                "reason": (
                    "배출계수가 설정되어 있지 않습니다."
                ),
            }

        try:
            numeric_value = float(value)
        except (TypeError, ValueError):
            return {
                "item": item_name,
                "status": "invalid_value",
                "carbon_kg": None,
                "reason": (
                    "사용량 또는 금액을 숫자로 "
                    "변환할 수 없습니다."
                ),
            }

        if numeric_value < 0:
            return {
                "item": item_name,
                "status": "invalid_value",
                "carbon_kg": None,
                "reason": (
                    "사용량 또는 금액은 "
                    "음수일 수 없습니다."
                ),
            }

        carbon_kg = numeric_value * factor

        return {
            "item": item_name,
            "status": "calculated",
            "input_value": numeric_value,
            "factor": factor,
            "factor_unit": coefficient["unit"],
            "carbon_basis": coefficient["carbon_basis"],
            "carbon_kg": round(carbon_kg, 6),
            "calculation_type": coefficient_type,
            "source": coefficient["source"],
        }

    def calculate_category(
        self,
        usage_item: str,
        usage_value: float | None = None,
        cost_item: str | None = None,
        cost_value: float | None = None,
    ) -> dict[str, Any]:
        """
        하나의 관리비 카테고리에 대해
        USAGE → COST 순서로 자동 계산한다.

        사용량이 존재하면 USAGE 방식을 우선 적용한다.

        사용량이 없거나 사용할 수 없는 경우
        COST 방식으로 fallback 한다.

        사용량과 금액이 모두 있어도 둘을 동시에 계산하지 않아
        탄소배출량 중복 계산을 방지한다.
        """
        # 1순위: 실제 사용량 기반 계산
        if usage_value is not None:
            usage_result = self.calculate_with_factor(
                usage_item,
                usage_value,
            )

            if usage_result.get("status") == "calculated":
                usage_result["fallback_used"] = False
                usage_result["priority_used"] = 1

                return usage_result

        # 2순위: 세부 요금 기반 fallback
        if cost_item is not None and cost_value is not None:
            cost_result = self.calculate_with_factor(
                cost_item,
                cost_value,
            )

            if cost_result.get("status") == "calculated":
                cost_result["fallback_used"] = True
                cost_result["priority_used"] = 2
                cost_result["fallback_from"] = usage_item

                return cost_result

        return {
            "item": usage_item,
            "status": "not_calculable",
            "carbon_kg": None,
            "fallback_used": False,
            "reason": (
                "사용량과 세부 요금 모두 "
                "탄소 계산에 사용할 수 없습니다."
            ),
        }

    def calculate_from_parsed_bill(
        self,
        parsed_bill: dict[str, Any],
    ) -> dict[str, Any]:
        """
        UtilityBillParser의 결과를 받아
        calculate_bill()에 자동으로 연결한다.

        Parser
            ↓
        calculate_from_parsed_bill()
            ↓
        calculate_bill()
            ↓
        탄소배출량 계산
        """
        if not parsed_bill:
            return {
                "status": "not_calculable",
                "items": [],
                "total_carbon_kg": 0.0,
                "warnings": [
                    "관리비 Parser 결과가 없습니다."
                ],
            }

        if parsed_bill.get("status") != "parsed":
            return {
                "status": "not_calculable",
                "items": [],
                "total_carbon_kg": 0.0,
                "warnings": (
                    parsed_bill.get("warnings")
                    or ["관리비 Parser 결과가 유효하지 않습니다."]
                ),
            }

        data = parsed_bill.get("data", {})

        electricity = data.get("electricity", {})
        gas = data.get("gas", {})
        water = data.get("water", {})
        district_heating = data.get(
            "district_heating",
            {},
        )
        hot_water = data.get("hot_water", {})
        sewer = data.get("sewer", {})

        # 지역난방 계수는 Gcal 기준이다.
        # Parser가 MWh를 반환하면 Gcal로 변환한다.
        district_heating_usage = district_heating.get("usage")
        district_heating_unit = district_heating.get("usage_unit")
        conversion_warnings: list[str] = []

        if district_heating_usage is not None:
            if district_heating_unit == "MWh":
                original_mwh = float(district_heating_usage)

                # 1 MWh = 3.6 GJ
                # 1 Gcal = 4.1868 GJ
                # 따라서 1 MWh ≈ 0.859845 Gcal
                district_heating_usage = (
                    original_mwh * 0.859845227858985
                )

                conversion_warnings.append(
                    f"지역난방 사용량 {original_mwh:g} MWh를 "
                    f"{district_heating_usage:.6f} Gcal로 "
                    "변환하여 계산했습니다."
                )

            elif district_heating_unit not in (None, "Gcal"):
                conversion_warnings.append(
                    "지원하지 않는 지역난방 사용량 단위"
                    f"({district_heating_unit})이므로 "
                    "지역난방 탄소 계산에서 제외했습니다."
                )

                district_heating_usage = None

        carbon_result = self.calculate_bill(
            electricity_usage=electricity.get("usage"),
            electricity_cost=electricity.get("cost"),
            gas_usage=gas.get("usage"),
            gas_cost=gas.get("cost"),
            water_usage=water.get("usage"),
            water_cost=water.get("cost"),
            district_heating_usage=district_heating_usage,
            hot_water_usage=hot_water.get("usage"),
            sewer_usage=sewer.get("usage"),
            management_fee_total=data.get(
                "management_fee_total"
            ),
        )

        # Parser 단계에서 발생한 warning도 최종 결과에 합친다.
        parser_warnings = parsed_bill.get(
            "warnings",
            [],
        )

        existing_warnings = carbon_result.get(
            "warnings",
            [],
        )

        carbon_result["warnings"] = list(
            dict.fromkeys(
                parser_warnings
                + conversion_warnings
                + existing_warnings
            )
        )

        return carbon_result

    def calculate_bill(
        self,
        electricity_usage: float | None = None,
        electricity_cost: float | None = None,
        gas_usage: float | None = None,
        gas_cost: float | None = None,
        water_usage: float | None = None,
        water_cost: float | None = None,
        district_heating_usage: float | None = None,
        hot_water_usage: float | None = None,
        sewer_usage: float | None = None,
        management_fee_total: float | None = None,
    ) -> dict[str, Any]:
        """
        관리비 고지서 전체의 탄소배출량을 계산한다.

        계산 순서
        ------------------------------------------------
        전기:
            전기 사용량(kWh)
            → 없으면 전기비(KRW)

        도시가스:
            도시가스 사용량(Nm³)
            → 없으면 도시가스비(KRW)

        수도:
            수도 사용량(m³)
            → 없으면 수도비(KRW)

        지역난방:
            열 사용량(Gcal)

        급탕:
            실제 열원 정보가 필요하므로 현재 직접 계산하지 않음

        하수:
            현재 UNSUPPORTED

        관리비 총액:
            위의 세부 항목 중 계산 가능한 항목이 하나도 없을 때만
            COST_TOTAL 방식으로 최종 fallback
        ------------------------------------------------

        세부 항목과 관리비 총액을 동시에 계산하지 않으므로
        중복 계산을 방지한다.
        """
        results: list[dict[str, Any]] = []
        warnings: list[str] = []

        # --------------------------------------------------
        # 전기
        # --------------------------------------------------
        electricity_result = self.calculate_category(
            usage_item="전기",
            usage_value=electricity_usage,
            cost_item="전기비",
            cost_value=electricity_cost,
        )

        if electricity_result.get("status") == "calculated":
            results.append(electricity_result)

        # --------------------------------------------------
        # 도시가스
        # --------------------------------------------------
        gas_result = self.calculate_category(
            usage_item="도시가스(LNG)",
            usage_value=gas_usage,
            cost_item="도시가스비",
            cost_value=gas_cost,
        )

        if gas_result.get("status") == "calculated":
            results.append(gas_result)

        # --------------------------------------------------
        # 수도
        # --------------------------------------------------
        water_result = self.calculate_category(
            usage_item="수도",
            usage_value=water_usage,
            cost_item="수도비",
            cost_value=water_cost,
        )

        if water_result.get("status") == "calculated":
            results.append(water_result)

        # --------------------------------------------------
        # 지역난방
        # --------------------------------------------------
        if district_heating_usage is not None:
            district_result = self.calculate_with_factor(
                "지역난방",
                district_heating_usage,
            )

            if district_result.get("status") == "calculated":
                district_result["fallback_used"] = False
                district_result["priority_used"] = 1

                results.append(district_result)

        # --------------------------------------------------
        # 급탕
        # --------------------------------------------------
        if hot_water_usage is not None:
            hot_water_result = self.calculate_with_factor(
                "급탕",
                hot_water_usage,
            )

            if hot_water_result.get("status") == "calculated":
                results.append(hot_water_result)

            else:
                warnings.append(
                    "급탕 사용량이 인식되었지만 실제 열원 "
                    "(도시가스/지역난방/전기) 정보가 없어 "
                    "탄소배출량 계산에서 제외했습니다."
                )

        # --------------------------------------------------
        # 하수
        # --------------------------------------------------
        if sewer_usage is not None:
            sewer_result = self.calculate_with_factor(
                "하수",
                sewer_usage,
            )

            if sewer_result.get("status") == "calculated":
                results.append(sewer_result)

            else:
                warnings.append(
                    "하수 사용량이 인식되었지만 현재 적용 가능한 "
                    "단일 배출계수가 없어 탄소배출량 계산에서 "
                    "제외했습니다."
                )

        # --------------------------------------------------
        # 세부 항목 계산 성공 여부
        # --------------------------------------------------
        detailed_calculation_used = len(results) > 0

        # --------------------------------------------------
        # 3순위: 관리비 총액 COST_TOTAL fallback
        #
        # 세부 항목을 하나도 계산하지 못한 경우에만 적용
        # --------------------------------------------------
        total_fallback_used = False

        if (
            not detailed_calculation_used
            and management_fee_total is not None
        ):
            total_result = self.calculate_with_factor(
                "주거관리비",
                management_fee_total,
            )

            if total_result.get("status") == "calculated":
                total_result["fallback_used"] = True
                total_result["priority_used"] = 3
                total_result["fallback_reason"] = (
                    "계산 가능한 세부 사용량 또는 "
                    "세부 요금이 없어 관리비 총액 계수를 "
                    "적용했습니다."
                )

                results.append(total_result)
                total_fallback_used = True

        # --------------------------------------------------
        # 최종 합계
        # --------------------------------------------------
        total_carbon_kg = round(
            sum(
                result["carbon_kg"]
                for result in results
                if result.get("carbon_kg") is not None
            ),
            6,
        )

        # --------------------------------------------------
        # 탄소 기준 확인
        #
        # 현재 Excel에는 CO₂eq와 CO₂ 계수가 혼재할 수 있으므로
        # 최종 합계를 무조건 kgCO₂eq라고 표기하지 않는다.
        # --------------------------------------------------
        carbon_bases = sorted(
            {
                result.get("carbon_basis")
                for result in results
                if result.get("carbon_basis")
            }
        )

        if len(carbon_bases) > 1:
            warnings.append(
                "계산 결과에 서로 다른 탄소 기준 "
                "(예: CO₂와 CO₂eq)이 포함되어 있습니다. "
                "각 항목의 carbon_basis를 확인하세요."
            )

        # --------------------------------------------------
        # 아무것도 계산하지 못한 경우
        # --------------------------------------------------
        if not results:
            warnings.append(
                "탄소배출량을 계산할 수 있는 "
                "사용량 또는 요금 정보가 없습니다."
            )

        return {
            "status": (
                "calculated"
                if results
                else "not_calculable"
            ),
            "items": results,
            "total_carbon_kg": total_carbon_kg,
            "carbon_bases": carbon_bases,
            "detailed_calculation_used": (
                detailed_calculation_used
            ),
            "total_fallback_used": total_fallback_used,
            "warnings": warnings,
        }


# 다른 파일에서 바로 사용할 수 있도록 기본 인스턴스 생성
utility_carbon_service = UtilityCarbonService()